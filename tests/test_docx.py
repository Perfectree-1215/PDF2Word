"""Word output: build a .docx from a small Markdown sample and read it back.

Run:  python -m pytest tests/test_docx.py -q      (or)      python tests/test_docx.py
Needs python-docx; native equations additionally need latex2mathml + Microsoft Office's
MML2OMML.XSL (the test accepts the LaTeX-text fallback when they are missing).
"""
from __future__ import annotations

import re
import sys
import tempfile
import unicodedata
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pdf2word.config import Options  # noqa: E402
from pdf2word.docx_out import OmmlConverter, build_docx  # noqa: E402
from pdf2word.metadata import PaperMeta  # noqa: E402
from pdf2word.obsidian import image_link  # noqa: E402

SAMPLE_MD = """# Inline Equation Images in Word Documents

Kwang Hee Kim, Jane Doe

> [!abstract] Abstract
> This note tests the Word writer with **bold**, _italic_ and $\\rho_{f}$ inline math.

## 1 Introduction

The particle _i_<sup>th</sup> velocity $u_i$ and the density ρ<sub>f</sub> are used [^1].

### <u>2.3.1 Steady Flamelet (SF) Model</u>

Some <u>underlined words</u> and a <mark>highlighted</mark> phrase in running text.

![[_asset-Kim-2024/fig1.png]]

**Figure 1.** A tiny test picture. ^fig-1

![[_asset-Kim-2024/eq-1.png]]
$$
m_{\\mathrm{p}}\\frac{\\mathrm{d}v}{\\mathrm{d}t}=\\sum F_{\\mathrm{c}}+m_{\\mathrm{p}}g \\tag{1}
$$
^eq-1

$$
\\begin{gathered}
a = b \\\\
c = d
\\end{gathered} \\tag{2}
$$
^eq-2

| Quantity | Value |
|---|---|
| Density | 1000 kg/m<sup>3</sup> |
| Viscosity | 0.001 Pa s |

- first item
- second item

%% equation text: hidden comment %%

## References

[^1]: A. Author, "A paper," *Journal*, 2024.
"""


def _make_png(path: Path, size=(120, 40), dpi=144):
    from PIL import Image
    img = Image.new("L", size, 255)
    for x in range(10, size[0] - 10):
        img.putpixel((x, size[1] // 2), 0)
    img.save(path, dpi=(dpi, dpi))


def test_build_docx():
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        assets = base / "_asset-Kim-2024"
        assets.mkdir()
        _make_png(assets / "fig1.png", size=(400, 200), dpi=200)
        _make_png(assets / "eq-1.png", size=(240, 60), dpi=144)
        meta = PaperMeta(title="Inline Equation Images in Word Documents",
                         authors=["Kwang Hee Kim", "Jane Doe"], first_author="Kim", year=2024,
                         journal="Test Journal", doi="10.1000/test", keywords=["a", "b"])
        stats = {}
        out = build_docx(SAMPLE_MD, meta, base, base / "out.docx", Options(), stats)
        assert out.exists() and out.stat().st_size > 5000

        import docx
        d = docx.Document(str(out))
        texts = [p.text for p in d.paragraphs]
        headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
        assert headings[0] == meta.title                       # title once, from metadata
        assert texts.count(meta.title) == 1                    # the '# Title' line is not repeated
        assert "Abstract" in headings and "1 Introduction" in headings and "References" in headings
        assert any("Keywords:" in t for t in texts)
        assert len(d.inline_shapes) == 2                       # figure + equation image
        assert len(d.tables) == 1 and len(d.tables[0].rows) == 3
        assert any(p.style.name == "Caption" for p in d.paragraphs)
        assert any(p.style.name == "List Bullet" for p in d.paragraphs)
        assert not any("hidden comment" in t for t in texts)   # %% comments are dropped
        xml = d.element.xml
        n_native = xml.count(f'<m:oMath ') + xml.count("<m:oMath>")
        if OmmlConverter().available:
            assert stats.get("docx_eq_native", 0) >= 3 and n_native >= 3   # 2 blocks + inline math
            assert "(1)" in "".join(texts) and "(2)" in "".join(texts)
        else:
            assert stats.get("docx_eq_text", 0) >= 2
            assert any("\\frac{\\mathrm{d}v}" in t for t in texts)
        # superscript/subscript runs survive
        sup = [r for p in d.paragraphs for r in p.runs if r.font.superscript]
        sub = [r for p in d.paragraphs for r in p.runs if r.font.subscript]
        assert sup and sub
        # <u>/<mark> become formatting, never literal tags; headings lose the tags entirely
        assert "<u>" not in xml and "&lt;u&gt;" not in xml
        assert "2.3.1 Steady Flamelet (SF) Model" in headings
        underlined = [r for p in d.paragraphs for r in p.runs if r.underline and "underlined words" in r.text]
        assert underlined


def test_prepare_latex():
    prep = OmmlConverter.prepare
    assert prep(r"\Bigl(a\Bigr) \tag{3}") == r"\left(a\right)"
    assert prep(r"\Big(x\Big)") == "(x)"
    assert prep(r"\begin{gathered}a\\b\end{gathered}") == r"\begin{matrix}a\\b\end{matrix}"
    assert prep(r"\left(a") == "(a"                         # unbalanced sizing commands dropped
    assert prep(r"\left(a\right)\rightarrow b") == r"\left(a\right)\rightarrow b"   # an arrow is no \right
    assert prep(r"\left. x\right|_{0}\right)") == r"x|_{0})"                       # no stray '.'


def _docx(tmp_path: Path, md: str, title: str = "Untitled Test", name: str = "out.docx"):
    """Build md into tmp_path/name under a title-only PaperMeta; return (document, stats)."""
    import docx
    stats = {}
    out = build_docx(md, PaperMeta(title=title), tmp_path, tmp_path / name, Options(), stats)
    return docx.Document(str(out)), stats


def _letters(text: str) -> Counter:
    """The letters and digits of text: what a reader would miss, ignoring markup and spacing."""
    return Counter(re.sub(r"[\W_]", "", text))


def test_prepare_big_operators_and_angle_delims():
    prep = OmmlConverter.prepare
    big_ops = ("cup", "cap", "oplus", "otimes", "odot", "sqcup", "uplus", "vee", "wedge",
               "circ", "star", "triangleup", "triangledown")
    for op in big_ops:                                  # \bigcup is an operator, not \big + "cup"
        assert prep(rf"\big{op}_{{i}} A_i") == rf"\big{op}_{{i}} A_i"
    assert prep(r"\bigl\langle x \bigr\rangle") == r"\left\langle x \right\rangle"
    assert prep(r"\Bigl\lfloor n \Bigr\rfloor") == r"\left\lfloor n \right\rfloor"
    assert prep(r"\bigl. \frac{a}{b} \bigr|_{0}") == r"\left. \frac{a}{b} \right|_{0}"
    assert prep(r"\Bigl. x |") == "x |"                 # unmatched null delimiter: no stray 'l' or '.'
    assert prep(r"\Bigm| y") == "| y"
    conv = OmmlConverter()
    if not conv.available:                              # no Office stylesheet: prepare() is checked above
        return
    def text(latex):                                    # what Word prints; an n-ary sign sits in <m:chr m:val>
        om = conv.omml(latex)
        signs = [v for e in om.iter() if e.tag.endswith("}chr") for v in e.attrib.values()]
        return "".join(signs) + "".join(om.itertext())

    for op, sym in zip(big_ops, "⋃⋂⨁⨂⨀⨆⨄⋁⋀◯★△▽"):
        t = text(rf"\big{op}_{{i=1}}^{{n}} A_i")
        assert sym in t and op not in t, (op, t)
    assert text(r"\bigl\langle x, y \bigr\rangle") == "⟨x,y⟩"
    assert "." not in text(r"\bigl. \frac{df}{dx} \bigr|_{x=0}")


def test_numbered_lines_keep_source_numbers(tmp_path):
    md = ("1. Heat the sample.\n2. Cool the sample.\n\nText between the lists.\n\n"
          "5. Fifth step of a list that starts at five.\n6. Sixth step.\n\n- a bullet item\n\n"
          "## References\n\n1. A. Author, First reference, 2020.\n2. B. Author, Second reference, 2021.\n")
    d, _ = _docx(tmp_path, md)
    texts = [p.text for p in d.paragraphs]
    for line in ("1. Heat the sample.", "2. Cool the sample.", "5. Fifth step of a list that starts at five.",
                 "6. Sixth step.", "1. A. Author, First reference, 2020.", "2. B. Author, Second reference, 2021."):
        assert line in texts                            # the number printed in the PDF, not Word's counter
    assert not any(p.style.name == "List Number" for p in d.paragraphs)
    assert [p.text for p in d.paragraphs if p.style.name == "List Bullet"] == ["a bullet item"]


def test_markdown_image_links_with_non_ascii_asset_dir(tmp_path):
    folders = ["_asset-김-2024",                                  # Hangul author name
               unicodedata.normalize("NFD", "_asset-김-2024"),     # ... spelled in jamo, as safe_name folds it
               "_asset-Manual-(v2),-rev'3",                        # ASCII that quote() escapes
               "_asset-100%-done"]                                  # a literal '%'
    cases = [(f, "markdown") for f in folders] + [("_asset-a%20b", "wiki")]   # '%20' is the real name
    for k, (folder, style) in enumerate(cases):
        base = tmp_path / f"case{k}"
        (base / folder).mkdir(parents=True)
        for name, size in (("fig1.png", (400, 200)), ("eq-1.png", (240, 60)), ("sym-1.png", (30, 30))):
            _make_png(base / folder / name, size=size)
        link = lambda name: image_link(f"{folder}/{name}", Path(name).stem, style)   # noqa: E731
        md = (f"{link('fig1.png')}\n\n**Figure 1.** A caption. ^fig-1\n\n{link('eq-1.png')}\n$$\na=b\n$$\n\n"
              f"A symbol {link('sym-1.png')} in running text.\n")
        d, stats = _docx(base, md)
        assert len(d.inline_shapes) == 3, (folder, style)            # figure, equation crop, symbol
        assert stats.get("docx_pictures") == 2 and stats.get("docx_inline_pictures") == 1


def test_non_latin_title_not_repeated(tmp_path):
    for k, title in enumerate(("한국어 합성 논문 제목의 반복 검사", "Синтетическая статья о выводе")):
        md = f"# {title}\n\nBody text.\n\n# 서론\n\nMore text.\n"
        d, _ = _docx(tmp_path, md, title=title, name=f"t{k}.docx")
        texts = [p.text for p in d.paragraphs]
        assert texts.count(title) == 1 and d.paragraphs[0].style.name == "Title"
        assert "서론" in texts                           # another level-1 heading is still written


def test_caret_at_line_end_is_not_block_id(tmp_path):
    md = ("A 32-bit counter overflows at 2^31\n\nThe plate area is 4 m^2\n\n"
          "**Figure 1.** A caption with an anchor. ^fig-1\n\nEquation follows.\n^eq-3\n")
    d, _ = _docx(tmp_path, md)
    texts = [p.text for p in d.paragraphs]
    assert "A 32-bit counter overflows at 2^31" in texts and "The plate area is 4 m^2" in texts
    assert "Figure 1. A caption with an anchor." in texts and "Equation follows." in texts
    assert not any("^fig-1" in t or "^eq-3" in t for t in texts)


def test_equation_image_sits_right_before_its_equation(tmp_path):
    (tmp_path / "_asset-x").mkdir()
    for name in ("eq-1.png", "eq-2.png"):
        _make_png(tmp_path / "_asset-x" / name, size=(240, 60))
    md = ("Before.\n\n![[_asset-x/eq-1.png]]\n$$\na=b \\tag{1}\n$$\n\n"
          "![[_asset-x/eq-2.png]]\n\n$$\nc=d \\tag{2}\n$$\n\n"
          "![[_asset-x/eq-9.png]]\n$$\ne=f \\tag{3}\n$$\n\nAfter.\n")      # eq-9.png does not exist
    d, _ = _docx(tmp_path, md)
    seq = ["pic" if "graphicData" in p._p.xml else p.text.rstrip()[-3:] if p.text.rstrip().endswith(")")
           else p.text for p in d.paragraphs[1:]]
    assert seq == ["Before.", "pic", "(1)", "pic", "(2)", "(3)", "After."]


def test_paragraph_text_is_preserved(tmp_path):
    title = "한국어 합성 논문 제목의 반복 검사"
    md = (f"# {title}\n\n1. Heat the sample.\n2. Cool the sample.\n\nText between the lists.\n\n"
          "5. Fifth step of a list that starts at five.\n6. Sixth step.\n\n- a bullet item\n\n"
          "A 32-bit counter overflows at 2^31\nThe plate area is 4 m^2\n\n"
          "**Figure 1.** A caption with an anchor. ^fig-1\n\n## References\n\n"
          "1. A. Author, First reference, 2020.\n2. B. Author, Second reference, 2021.\n")
    d, _ = _docx(tmp_path, md, title=title)
    # every letter and digit of the Markdown reaches Word exactly once (the '# title' line becomes the
    # Title paragraph); only the ^fig-1 anchor is markup
    assert _letters("".join(p.text for p in d.paragraphs)) == _letters(md.replace("^fig-1", ""))


if __name__ == "__main__":
    failures = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                if fn.__code__.co_argcount:
                    with tempfile.TemporaryDirectory() as tmp:
                        fn(Path(tmp))
                else:
                    fn()
                print("ok  ", name)
            except AssertionError as exc:
                failures += 1
                print("FAIL", name, exc)
    raise SystemExit(1 if failures else 0)
