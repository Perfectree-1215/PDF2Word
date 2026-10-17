"""Word (.docx) output: renders the Markdown produced by this tool into a Word document.

Equations are written as the cropped image of the formula followed by the OCR result,
either as a native Word equation (LaTeX -> MathML -> OMML through Microsoft Office's
MML2OMML.XSL stylesheet, when Office is installed) or as LaTeX text.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Optional

from .utils import log, strip_emphasis

# --------------------------------------------------------------------------- constants

_OFFICE_XSL = [
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\root\Office15\MML2OMML.XSL",
    "/Applications/Microsoft Word.app/Contents/Resources/MML2OMML.XSL",
]
_M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
CAPTION_RE = re.compile(r"^(Fig(?:ure)?|Table|Scheme|Algorithm|Chart|Plate)\.?\s*[0-9]+[A-Za-z]?\b", re.I)
_IMG_WIKI = re.compile(r"^!\[\[([^\]|]+)(?:\|[^\]]*)?\]\]\s*$")
_IMG_MD = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)\s*$")
_BLOCK_ID = re.compile(r"\s*\^[A-Za-z0-9-]+\s*$")
_TOKEN = re.compile(
    r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<![A-Za-z0-9])_(?!\s)[^_<>|\n]+?_(?![A-Za-z0-9])|\*(?!\s)[^*<>|\n]+?\*"
    r"|<sup>.*?</sup>|<sub>.*?</sub>|<u>.*?</u>|<mark>.*?</mark>|<(?:i|em)>.*?</(?:i|em)>|<(?:b|strong)>.*?</(?:b|strong)>"
    r"|`[^`]+`|\$[^$\n]+\$|\[\^[^\]]+\])", re.S)
_STRAY_TAG = re.compile(r"</?(?:u|mark|i|em|b|strong|s|del|span|font)[^>]*>")


# --------------------------------------------------------------------------- LaTeX -> OMML

class OmmlConverter:
    """LaTeX -> MathML (latex2mathml) -> OMML (Office's MML2OMML.XSL). `available` is False
    when either piece is missing; callers then fall back to LaTeX text."""

    def __init__(self, xsl_path: Optional[str] = None):
        self.available = False
        self.transform = None
        self._convert = None
        try:
            from latex2mathml.converter import convert as _convert
            from lxml import etree
        except Exception as exc:                       # optional dependencies
            log.debug("OMML unavailable: %s", exc)
            return
        candidates = [xsl_path] if xsl_path else _OFFICE_XSL
        for cand in candidates:
            if cand and Path(cand).exists():
                try:
                    self.transform = etree.XSLT(etree.parse(str(cand)))
                    self._convert = _convert
                    self._etree = etree
                    self.available = True
                    log.debug("OMML stylesheet: %s", cand)
                    break
                except Exception as exc:
                    log.debug("cannot load %s: %s", cand, exc)

    @staticmethod
    def prepare(latex: str) -> str:
        """Rewrite constructs latex2mathml does not know into ones it does."""
        s = re.sub(r"\\tag\*?\{[^}]*\}", "", latex)
        # \Bigl( ... \Bigr) -> \left( ... \right);  bare \Big( / \Big) -> plain delimiters
        s = re.sub(r"\\[Bb]igg?l\s*(\\?[(\[|{])", r"\\left\1", s)
        s = re.sub(r"\\[Bb]igg?r\s*(\\?[)\]|}])", r"\\right\1", s)
        s = re.sub(r"\{\\[Bb]igg?\s*/\}", "/", s)
        s = re.sub(r"\\[Bb]igg?m?\s*", "", s)
        s = s.replace(r"\begin{gathered}", r"\begin{matrix}").replace(r"\end{gathered}", r"\end{matrix}")
        s = s.replace(r"\begin{aligned}", r"\begin{matrix}").replace(r"\end{aligned}", r"\end{matrix}")
        s = re.sub(r"\\(?:displaystyle|textstyle|scriptstyle|nolimits|limits)\b", "", s)
        s = re.sub(r"\\stackrel\{\\rightarrow\}\{([^{}]*)\}", r"\\vec{\1}", s)
        s = re.sub(r"\\stackrel\{\\rightarrow\}(\\[a-zA-Z]+\{[^{}]*\}|\\?[A-Za-z0-9])", r"\\vec{\1}", s)
        s = re.sub(r"\\mathrm\{([^{}]*)\}", lambda m: "\\mathrm{" + m.group(1).replace(" ", "") + "}", s)
        s = s.replace("~", r"\;")
        if s.count(r"\left") != s.count(r"\right"):
            s = re.sub(r"\\(left|right)(?=[\s(\[|.{\\)\]])", "", s)
        return s.strip()

    def omml(self, latex: str, block: bool = True):
        """Return an lxml <m:oMath> element, or None when the conversion fails."""
        if not self.available:
            return None
        try:
            mathml = self._convert(self.prepare(latex), display="block" if block else "inline")
            tree = self._etree.fromstring(mathml.encode("utf-8"))
            result = self.transform(tree)
            root = result.getroot()
            if root is None:
                return None
            omath = root if root.tag == f"{{{_M_NS}}}oMath" else root.find(f".//{{{_M_NS}}}oMath")
            if omath is None:
                return None
            # detach from the oMathPara wrapper so it can live inside an ordinary paragraph
            parent = omath.getparent()
            if parent is not None:
                parent.remove(omath)
            return omath
        except Exception as exc:
            log.debug("OMML conversion failed for %r: %s: %s", latex[:60], type(exc).__name__, exc)
            return None


# --------------------------------------------------------------------------- document builder

class DocxBuilder:
    def __init__(self, meta, base_dir: Path, options, stats=None):
        import docx  # python-docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.shared import Cm, Pt, RGBColor
        self._docx = docx
        self._ALIGN = WD_ALIGN_PARAGRAPH
        self._TAB = WD_TAB_ALIGNMENT
        self._Cm, self._Pt, self._RGB = Cm, Pt, RGBColor
        self.meta = meta
        self.base_dir = Path(base_dir)
        self.opt = options
        self.stats = stats if stats is not None else {}
        self.omml = OmmlConverter() if getattr(options, "docx_math", "native") == "native" else None
        self.doc = docx.Document()
        self._page_setup()
        self.usable_cm = 17.0

    # ---- page / styles ----------------------------------------------------
    def _page_setup(self):
        Cm = self._Cm
        for section in self.doc.sections:
            section.page_width, section.page_height = Cm(21.0), Cm(29.7)     # A4
            section.left_margin = section.right_margin = Cm(2.0)
            section.top_margin = section.bottom_margin = Cm(2.0)
        normal = self.doc.styles["Normal"]
        normal.font.size = self._Pt(10.5)
        normal.paragraph_format.space_after = self._Pt(6)

    # ---- entry point ------------------------------------------------------
    def build(self, md_body: str, out_path: Path) -> Path:
        # the paper's own "Keywords:" paragraph makes the metadata line redundant
        self._front_matter(keywords=not re.search(r"(?im)^[*_ ]*keywords?[*_ ]*[:：]", md_body))
        self._render_blocks(md_body)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out_path))
        return out_path

    def _front_matter(self, keywords: bool = True):
        m = self.meta
        self.doc.add_heading(m.title or "Untitled", 0)
        container = getattr(m, "container", "")
        if container:                      # one chapter / page group of a bigger document
            pages = getattr(m, "part_pages", "")
            self.doc.add_paragraph(container + (f"  ·  pp. {pages}" if pages else ""))
        if m.authors:
            self.doc.add_paragraph(", ".join(m.authors))
        bits = []
        if m.journal:
            bits.append(m.journal)
        if m.year:
            bits.append(str(m.year))
        if m.volume:
            bits.append(f"vol. {m.volume}" + (f"({m.issue})" if m.issue else ""))
        if m.pages:
            bits.append(f"pp. {m.pages}")
        if bits:
            self.doc.add_paragraph(", ".join(bits))
        if m.doi or m.url:
            self.doc.add_paragraph(f"DOI: {m.doi}" if m.doi else m.url)
        if m.keywords and keywords:
            p = self.doc.add_paragraph()
            p.add_run("Keywords: ").bold = True
            p.add_run("; ".join(m.keywords))

    # ---- block level --------------------------------------------------------
    def _render_blocks(self, md: str):
        lines = md.split("\n")
        i, n = 0, len(lines)
        para: list = []

        def flush():
            if para:
                text = " ".join(s.strip() for s in para).strip()
                if text:
                    self._paragraph(text)
                para.clear()

        pending_eq_image: Optional[Path] = None
        while i < n:
            line = lines[i]
            stripped = line.strip()
            if not stripped:
                flush()
                i += 1
                continue
            if stripped.startswith("%%") or stripped.startswith("<!--"):
                i += 1
                continue
            if _BLOCK_ID.fullmatch(" " + stripped) or re.fullmatch(r"\^[A-Za-z0-9-]+", stripped):
                i += 1
                continue
            m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if m:
                flush()
                level, text = len(m.group(1)), strip_emphasis(m.group(2))
                if level == 1 and self._same_title(text):
                    i += 1
                    continue
                self.doc.add_heading(text, min(max(level - 1, 1), 6))
                i += 1
                continue
            if stripped.startswith(">"):
                flush()
                i = self._blockquote(lines, i)
                continue
            if stripped == "$$":
                flush()
                i = self._equation(lines, i, pending_eq_image)
                pending_eq_image = None
                continue
            img = _IMG_WIKI.match(stripped) or _IMG_MD.match(stripped)
            if img:
                flush()
                path = self._resolve(img.group(1))
                # an equation image is followed by its $$ block: keep them together
                nxt = next((ln.strip() for ln in lines[i + 1:i + 3] if ln.strip()), "")
                if path and path.stem.startswith("eq-") and nxt == "$$":
                    pending_eq_image = path
                else:
                    self._picture(path, equation=bool(path and path.stem.startswith("eq-")))
                i += 1
                continue
            if stripped.startswith("|"):
                flush()
                i = self._table(lines, i)
                continue
            m = re.match(r"^([-*+]|\d+[.)])\s+(.*)$", stripped)
            if m and not para:
                flush()
                style = "List Number" if m.group(1)[0].isdigit() else "List Bullet"
                self._paragraph(m.group(2), style=style)
                i += 1
                continue
            if stripped.startswith("---") and set(stripped) <= {"-", " "}:
                i += 1
                continue
            if re.search(r"\s\^(fig|table|scheme|algorithm|chart|plate)-[\w-]+\s*$", line):
                para.append("\x00CAPTION\x00" + _BLOCK_ID.sub("", line))
            else:
                para.append(_BLOCK_ID.sub("", line))
            i += 1
        flush()

    def _same_title(self, text: str) -> bool:
        norm = lambda s: re.sub(r"[^a-z0-9]+", "", s.lower())
        t = norm(self.meta.title or "")
        return bool(t) and (norm(text) == t or norm(text)[:40] == t[:40])

    def _blockquote(self, lines: list, i: int) -> int:
        body = []
        title = None
        while i < len(lines) and lines[i].strip().startswith(">"):
            content = lines[i].strip()[1:].strip()
            m = re.match(r"^\[!(\w+)\]\s*(.*)$", content)
            if m:
                title = m.group(2) or m.group(1).capitalize()
            else:
                body.append(content)
            i += 1
        if title:
            self.doc.add_heading(title, 1)
        text = []
        for ln in body:
            if ln:
                text.append(ln)
            elif text:
                self._paragraph(" ".join(text), style="Quote" if not title else None)
                text = []
        if text:
            self._paragraph(" ".join(text), style="Quote" if not title else None)
        return i

    def _table(self, lines: list, i: int) -> int:
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                rows.append(cells)
            i += 1
        if not rows:
            return i
        ncols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        try:
            table.style = self.doc.styles["Table Grid"]
        except KeyError:
            pass
        for r, cells in enumerate(rows):
            for c in range(ncols):
                cell = table.cell(r, c)
                cell.text = ""
                self._runs(cell.paragraphs[0], cells[c] if c < len(cells) else "", bold=(r == 0))
        self.doc.add_paragraph()
        return i

    # ---- images & equations --------------------------------------------------
    def _resolve(self, ref: str) -> Optional[Path]:
        ref = ref.strip().replace("\\", "/")
        for cand in (self.base_dir / ref, self.base_dir.parent / ref, Path(ref)):
            if cand.exists():
                return cand
        log.warning("  docx: image not found: %s", ref)
        return None

    def _image_width_cm(self, path: Path, max_cm: float) -> float:
        try:
            from PIL import Image
            with Image.open(path) as im:
                px_w = im.width
                dpi = im.info.get("dpi", (0, 0))[0] or 150.0
        except Exception:
            return max_cm
        width_cm = px_w / float(dpi) * 2.54
        return min(width_cm, max_cm)

    def _picture(self, path: Optional[Path], equation: bool = False):
        if path is None:
            return
        p = self.doc.add_paragraph()
        p.alignment = self._ALIGN.CENTER
        width = self._image_width_cm(path, 14.0 if equation else self.usable_cm)
        try:
            p.add_run().add_picture(str(path), width=self._Cm(width))
            self.stats["docx_pictures"] = self.stats.get("docx_pictures", 0) + 1
        except Exception as exc:
            log.warning("  docx: cannot embed %s (%s)", path.name, exc)

    def _equation(self, lines: list, i: int, image: Optional[Path]) -> int:
        i += 1
        body = []
        while i < len(lines) and lines[i].strip() != "$$":
            body.append(lines[i])
            i += 1
        i += 1                                         # closing $$
        latex = "\n".join(body).strip()
        label = None
        m = re.search(r"\\tag\*?\{([^}]*)\}", latex)
        if m:
            label = m.group(1)
            latex = (latex[:m.start()] + latex[m.end():]).strip()
        if image is not None:
            self._picture(image, equation=True)
        self._equation_paragraph(latex, label)
        return i

    def _equation_paragraph(self, latex: str, label: Optional[str]):
        p = self.doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.tab_stops.add_tab_stop(self._Cm(self.usable_cm / 2), self._TAB.CENTER)
        fmt.tab_stops.add_tab_stop(self._Cm(self.usable_cm), self._TAB.RIGHT)
        p.add_run("\t")
        omath = self.omml.omml(latex, block=True) if self.omml else None
        if omath is not None:
            p._p.append(omath)
            self.stats["docx_eq_native"] = self.stats.get("docx_eq_native", 0) + 1
        else:
            run = p.add_run(latex.replace("\n", " "))
            run.font.name = "Cambria Math"
            run.font.size = self._Pt(9.5)
            self.stats["docx_eq_text"] = self.stats.get("docx_eq_text", 0) + 1
        if label:
            p.add_run(f"\t({label})")

    # ---- inline level ----------------------------------------------------------
    def _paragraph(self, text: str, style: Optional[str] = None):
        # A caption is a short paragraph whose "Figure 3." label is emphasised or carries a
        # ^fig-3 anchor; "Figure 1 shows ..." in running text is not one.
        is_caption = "\x00CAPTION\x00" in text
        text = text.replace("\x00CAPTION\x00", "")
        if not is_caption and len(text) < 700:
            is_caption = bool(re.match(r"^(\*\*|\*|_)\s*(Fig(?:ure)?|Table|Scheme|Algorithm|Chart|Plate)\.?\s*\d+", text))
        if style is None and is_caption:
            style = "Caption"
        try:
            p = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        except KeyError:
            p = self.doc.add_paragraph()
        self._runs(p, text)
        if is_caption:
            p.alignment = self._ALIGN.LEFT
        if text.lstrip("*_ ").startswith("⚠"):            # OCR warning under an equation
            p.alignment = self._ALIGN.CENTER
            for r in p.runs:
                r.font.size = self._Pt(8.5)
                r.font.color.rgb = self._RGB(0xB0, 0x60, 0x60)
        return p

    def _runs(self, p, text: str, bold: bool = False, italic: bool = False,
              underline: bool = False, highlight: bool = False):
        pos = 0
        fmt = dict(bold=bold, italic=italic, underline=underline, highlight=highlight)
        for m in _TOKEN.finditer(text):
            if m.start() > pos:
                self._run(p, text[pos:m.start()], **fmt)
            tok = m.group(0)
            if tok.startswith("***"):
                self._run(p, tok[3:-3], **dict(fmt, bold=True, italic=True))
            elif tok.startswith("**"):
                self._runs(p, tok[2:-2], **dict(fmt, bold=True))
            elif tok.startswith("<sup>"):
                self._run(p, strip_emphasis(tok[5:-6]), sup=True, **fmt)
            elif tok.startswith("<sub>"):
                self._run(p, strip_emphasis(tok[5:-6]), sub=True, **fmt)
            elif tok.startswith("<u>"):                # underlined text (pymupdf4llm emits <u>)
                self._runs(p, tok[3:-4], **dict(fmt, underline=True))
            elif tok.startswith("<mark>"):             # highlighted text
                self._runs(p, tok[6:-7], **dict(fmt, highlight=True))
            elif tok.startswith("<i>") or tok.startswith("<em>"):
                self._runs(p, re.sub(r"^<(?:i|em)>|</(?:i|em)>$", "", tok), **dict(fmt, italic=True))
            elif tok.startswith("<b>") or tok.startswith("<strong>"):
                self._runs(p, re.sub(r"^<(?:b|strong)>|</(?:b|strong)>$", "", tok), **dict(fmt, bold=True))
            elif tok.startswith("`"):
                r = self._run(p, tok[1:-1], **fmt)
                r.font.name = "Consolas"
            elif tok.startswith("$"):
                self._inline_math(p, tok[1:-1])
            elif tok.startswith("[^"):
                self._run(p, "[" + tok[2:-1] + "]", sup=True, **fmt)
            else:                                      # *italic* / _italic_
                self._runs(p, tok[1:-1], **dict(fmt, italic=True))
            pos = m.end()
        if pos < len(text):
            self._run(p, text[pos:], **fmt)

    def _run(self, p, text: str, bold=False, italic=False, underline=False, highlight=False,
             sup=False, sub=False):
        text = _STRAY_TAG.sub("", text.replace("<br>", "\n"))   # unmatched tags never show literally
        r = p.add_run(text)
        r.bold = bold or None
        r.italic = italic or None
        if underline:
            r.underline = True
        if highlight:
            from docx.enum.text import WD_COLOR_INDEX
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if sup:
            r.font.superscript = True
        if sub:
            r.font.subscript = True
        return r

    def _inline_math(self, p, latex: str):
        omath = self.omml.omml(latex, block=False) if self.omml else None
        if omath is not None:
            p._p.append(omath)
            self.stats["docx_eq_native"] = self.stats.get("docx_eq_native", 0) + 1
            return
        r = p.add_run(latex)
        r.font.name = "Cambria Math"
        r.italic = True


def build_docx(md_body: str, meta, base_dir: Path, out_path: Path, options, stats=None) -> Path:
    """Render the (frontmatter-free) Markdown body into out_path. Images are resolved
    relative to base_dir (the folder holding the _asset-* directory)."""
    builder = DocxBuilder(meta, base_dir, options, stats)
    return builder.build(md_body, Path(out_path))
